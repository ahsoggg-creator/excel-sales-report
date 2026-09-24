import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

COLUMNS = {
    "дата": "Дата",
    "товар": "Товар",
    "количество": "Количество",
    "кол-во": "Количество",
    "кол-во, шт": "Количество",
    "цена": "Цена",
    "менеджер": "Менеджер",
}
DATE_FORMATS = ["%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"]
MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]


def parse_date(value):
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def parse_number(value):
    if isinstance(value, (int, float)):
        return value
    text = re.sub(r"[^\d,.]", "", str(value)).replace(",", ".")
    return float(text) if text else None


def read_report(path):
    raw = pd.read_excel(path, header=None)
    header_row = None
    for i, row in raw.iterrows():
        if any(str(cell).strip().lower() == "дата" for cell in row):
            header_row = i
            break
    if header_row is None:
        print(f"  {path.name}: не нашел строку с заголовками, пропускаю")
        return None

    df = raw.iloc[header_row + 1:].copy()
    df.columns = [COLUMNS.get(str(c).strip().lower(), str(c).strip()) for c in raw.iloc[header_row]]
    df = df[[c for c in dict.fromkeys(COLUMNS.values()) if c in df.columns]].copy()
    df["Филиал"] = path.stem
    return df


def clean(df):
    df = df.dropna(how="all", subset=["Дата", "Товар", "Количество", "Цена"])
    df = df[~df["Дата"].astype(str).str.strip().str.lower().str.startswith("итого")].copy()

    df["Дата"] = df["Дата"].map(parse_date)
    df["Товар"] = df["Товар"].astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    df["Менеджер"] = df["Менеджер"].astype(str).str.strip()
    df["Количество"] = df["Количество"].map(parse_number)
    df["Цена"] = df["Цена"].map(parse_number)

    df = df.dropna(subset=["Дата", "Количество", "Цена"])
    df = df.drop_duplicates()
    df["Количество"] = df["Количество"].astype(int)
    df["Сумма"] = df["Количество"] * df["Цена"]
    return df.sort_values(["Дата", "Филиал"]).reset_index(drop=True)


def make_pivot(df):
    df = df.copy()
    df["Месяц"] = df["Дата"].dt.to_period("M")
    pivot = df.pivot_table(index="Филиал", columns="Месяц", values="Сумма", aggfunc="sum", fill_value=0)
    pivot.columns = [f"{MONTHS[p.month - 1]} {p.year}" for p in pivot.columns]
    pivot["Итого"] = pivot.sum(axis=1)
    pivot = pivot.sort_values("Итого", ascending=False)
    pivot.loc["Всего"] = pivot.sum()
    return pivot


def style_sheet(ws, money_cols=(), date_cols=()):
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for col in ws.columns:
        letter = col[0].column_letter
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = min(width + 3, 40)
        header = col[0].value
        for cell in col[1:]:
            if header in money_cols:
                cell.number_format = "#,##0"
            elif header in date_cols:
                cell.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"


def add_chart(ws, rows, month_count):
    chart = BarChart()
    chart.title = "Выручка по филиалам по месяцам"
    chart.y_axis.title = "руб."
    chart.height = 10
    chart.width = 22

    data = Reference(ws, min_col=2, max_col=month_count + 1, min_row=1, max_row=rows + 1)
    branches = Reference(ws, min_col=1, min_row=2, max_row=rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(branches)
    ws.add_chart(chart, f"A{rows + 5}")


def main():
    parser = argparse.ArgumentParser(description="Сборка отчетов филиалов в один файл")
    parser.add_argument("--input", default="input", help="папка с отчетами .xlsx")
    parser.add_argument("--output", default="report.xlsx", help="куда сохранить итог")
    args = parser.parse_args()

    files = sorted(Path(args.input).glob("*.xlsx"))
    files = [f for f in files if not f.name.startswith("~$")]
    if not files:
        print(f"В папке {args.input} нет файлов .xlsx")
        return

    print(f"Найдено файлов: {len(files)}")
    frames = []
    for path in files:
        df = read_report(path)
        if df is not None:
            print(f"  {path.name}: {len(df)} строк")
            frames.append(df)

    raw = pd.concat(frames, ignore_index=True)
    data = clean(raw)
    pivot = make_pivot(data)
    print(f"Всего строк: {len(raw)}, после очистки: {len(data)} (убрано {len(raw) - len(data)})")

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        pivot.to_excel(writer, sheet_name="Сводная", index_label="Филиал")
        data.to_excel(writer, sheet_name="Данные", index=False)

    wb = load_workbook(args.output)
    style_sheet(wb["Данные"], money_cols=("Цена", "Сумма"), date_cols=("Дата",))
    wb["Данные"].auto_filter.ref = wb["Данные"].dimensions
    style_sheet(wb["Сводная"], money_cols=list(pivot.columns))
    for cell in wb["Сводная"][wb["Сводная"].max_row]:
        cell.font = Font(bold=True)
    add_chart(wb["Сводная"], rows=len(pivot) - 1, month_count=len(pivot.columns) - 1)
    wb.save(args.output)

    print(f"Готово: {args.output}")


if __name__ == "__main__":
    main()
